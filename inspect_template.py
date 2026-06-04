#!/usr/bin/env python3
"""
inspect_template.py — 分析 PPTX 模板结构，输出人类可读报告并可选生成 JSON 配置。

Usage:
    python inspect_template.py template.pptx
    python inspect_template.py template.pptx --json out.json
"""

import sys
import json
import argparse
from pathlib import Path
from pptx import Presentation

def emu_to_in(emu):
    return round(emu / 914400, 4)

def _looks_like_cover(slide):
    """True if slide appears to be a cover rather than a normal content slide."""
    for shape in slide.shapes:
        if not hasattr(shape, "text"):
            continue
        txt = shape.text.strip()
        if not txt:
            continue
        h = shape.height / 914400
        t = shape.top / 914400
        w = shape.width / 914400
        if h > 1.5 and w > 4 and t >= 0.8 and len(txt) > 40:
            return False
    return True

def _choose_content_slide(prs):
    """Mirror md_to_pptx.py: slide 1 is cover, slide 2 is content when applicable."""
    if len(prs.slides) >= 2 and _looks_like_cover(prs.slides[0]):
        return 1, prs.slides[1], True
    return 0, prs.slides[0], False

def inspect(template_path: str, json_out: str = None):
    prs = Presentation(template_path)

    sw = emu_to_in(prs.slide_width)
    sh = emu_to_in(prs.slide_height)

    content_idx, slide, has_cover = _choose_content_slide(prs)

    print(f"\n{'='*60}")
    print(f"模板：{template_path}")
    print(f"尺寸：{sw}\" × {sh}\"")
    print(f"总页数：{len(prs.slides)}")
    print(f"识别：{'封面页 + 正文页' if has_cover else '单一正文版式'}")
    print(f"用于检测正文版式的页面：第 {content_idx + 1} 页")
    if len(prs.slides) >= 3:
        print("提示：最后一页通常作为结束页/感谢页，不参与正文区域检测。")
    print(f"{'='*60}\n")

    print(f"{'ID':<5} {'名称':<22} {'类型':<18} {'左':<7} {'上':<7} {'宽':<7} {'高':<7}  文字预览")
    print("-" * 95)

    for s in slide.shapes:
        l = emu_to_in(s.left); t = emu_to_in(s.top)
        w = emu_to_in(s.width); h = emu_to_in(s.height)
        txt = (s.text.replace("\n"," ")[:55] if hasattr(s,"text") else "")
        print(f"{s.shape_id:<5} {s.name[:20]:<22} {str(s.shape_type)[:16]:<18} "
              f"{l:<7} {t:<7} {w:<7} {h:<7}  {txt}")

    # ── Auto-detect title and content boxes ──────────────────────────
    title_shape   = None
    content_shape = None
    deco_shapes   = []

    for s in slide.shapes:
        has_text = hasattr(s, "text")
        l = emu_to_in(s.left); t = emu_to_in(s.top)
        w = emu_to_in(s.width); h = emu_to_in(s.height)
        txt = s.text.strip() if has_text else ""

        # Title: wide, near top, not too tall.
        if has_text and t < 1.5 and h < 1.5 and w > 4:
            if title_shape is None or t < emu_to_in(title_shape.top):
                title_shape = s
            continue

        # Content: wide, lower, tall enough. Prefer the largest text area.
        if has_text and t >= 0.8 and h > 1.5 and w > 4:
            if content_shape is None:
                content_shape = s
            else:
                cur_area = emu_to_in(content_shape.width) * emu_to_in(content_shape.height)
                new_area = w * h
                if new_area > cur_area:
                    content_shape = s
            continue

        deco_shapes.append(s)

    # Slide number: near bottom and text-bearing.
    slide_num_shape = None
    for s in deco_shapes:
        if hasattr(s, "text") and emu_to_in(s.top) > sh * 0.8:
            slide_num_shape = s
            break

    print(f"\n{'─'*60}")
    print("自动识别结果：")
    print(f"  标题框  → {'Shape ' + str(title_shape.shape_id) + ' ' + repr(title_shape.name) if title_shape else '⚠️  未识别'}")
    print(f"  内容框  → {'Shape ' + str(content_shape.shape_id) + ' ' + repr(content_shape.name) if content_shape else '⚠️  未识别'}")
    print(f"  页码框  → {'Shape ' + str(slide_num_shape.shape_id) + ' ' + repr(slide_num_shape.name) if slide_num_shape else '⚠️  未识别'}")
    print(f"  装饰元素→ {len(deco_shapes)} 个")

    cfg = {
        "_note": "此文件由 inspect_template.py 自动生成，可手动修改后重新运行转换。",
        "slide_w": sw,
        "slide_h": sh,
    }

    if title_shape:
        cfg.update({
            "title_l": emu_to_in(title_shape.left),
            "title_t": emu_to_in(title_shape.top),
            "title_w": emu_to_in(title_shape.width),
            "title_h": emu_to_in(title_shape.height),
        })
    else:
        print("\n  ⚠️  未能自动识别标题框，请在生成的 JSON 中手动填写 title_l/t/w/h")

    if content_shape:
        cfg.update({
            "content_l": emu_to_in(content_shape.left),
            "content_t": emu_to_in(content_shape.top),
            "content_w": emu_to_in(content_shape.width),
            "content_h": emu_to_in(content_shape.height),
        })
    else:
        print("  ⚠️  未能自动识别内容框，请在生成的 JSON 中手动填写 content_l/t/w/h")

    if slide_num_shape:
        cfg.update({
            "slide_num_l": emu_to_in(slide_num_shape.left),
            "slide_num_t": emu_to_in(slide_num_shape.top),
            "slide_num_w": emu_to_in(slide_num_shape.width),
            "slide_num_h": emu_to_in(slide_num_shape.height),
        })

    cfg.update({
        "copy_deco": True,
        "font_main": "Microsoft YaHei",
        "title_color": "00B0F0",
        "h3_color": "005A87",
    })

    print(f"\n{'─'*60}")
    print("生成的 JSON 配置：")
    print(json.dumps(cfg, ensure_ascii=False, indent=2))

    if json_out:
        Path(json_out).write_text(
            json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"\n✅ 配置已写入：{json_out}")

    print(f"\n{'='*60}")
    print("提示：如识别有误，直接编辑 JSON 文件中的数值再重新运行转换即可。")
    print(f"{'='*60}\n")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="分析 PPTX 模板结构")
    parser.add_argument("template", help="PPTX 模板文件路径")
    parser.add_argument("--json", "-j", default=None,
                        help="将配置输出到此 JSON 文件路径（可选）")
    args = parser.parse_args()
    inspect(args.template, json_out=args.json)
